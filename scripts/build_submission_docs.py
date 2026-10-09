"""Render the Markdown submission reports to editable Word documents."""
from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def set_cell_text(cell, value, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(value.strip())
    run.bold = bold
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor(*color)


def build(source_name, output_name, title):
    source = ROOT / 'docs' / source_name
    lines = source.read_text(encoding='utf-8').splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.68)
    sec.bottom_margin = Inches(.68)
    sec.left_margin = Inches(.78)
    sec.right_margin = Inches(.78)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(31, 45, 64)
    normal.paragraph_format.space_after = Pt(3)
    for list_style_name in ('List Bullet', 'List Number'):
        list_style = doc.styles[list_style_name]
        list_style.font.name = 'Aptos'
        list_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        list_style.font.size = Pt(9)
        list_style.paragraph_format.space_after = Pt(2)
    for style_name, size, color in [('Heading 1', 19, (18, 51, 89)), ('Heading 2', 14, (24, 85, 129)), ('Heading 3', 11, (38, 113, 154))]:
        style = doc.styles[style_name]
        style.font.name = 'Aptos Display'
        style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    run.bold = True
    run.font.name = 'Aptos Display'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor(18, 51, 89)
    p = doc.add_paragraph('VeraLane · 2026 FinTech Hackathon 深圳国际金融科技大赛')
    p.paragraph_format.space_after = Pt(13)
    p.runs[0].font.size = Pt(10)
    p.runs[0].font.color.rgb = RGBColor(87, 108, 131)

    in_code = False
    code_lines = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if stripped.startswith('```'):
            if not in_code:
                in_code, code_lines = True, []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(.15)
                p.paragraph_format.right_indent = Inches(.12)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(7)
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), 'F1F5F9'); pPr.append(shd)
                r = p.add_run('\n'.join(code_lines))
                r.font.name = 'Consolas'; r.font.size = Pt(8)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped or stripped == '---':
            i += 1
            continue
        if stripped.startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{3,}:?', c.replace(' ', '')) for c in row):
                    table_lines.append(row)
                i += 1
            if table_lines:
                cols = max(map(len, table_lines))
                table = doc.add_table(rows=len(table_lines), cols=cols)
                table.style = 'Table Grid'
                table.autofit = True
                for ri, row in enumerate(table_lines):
                    for ci in range(cols):
                        val = row[ci] if ci < len(row) else ''
                        set_cell_text(table.cell(ri, ci), re.sub(r'[`*]', '', val), bold=ri == 0, color=(255,255,255) if ri == 0 else None)
                        if ri == 0: shade(table.cell(ri, ci), '183B60')
                        elif ri % 2 == 0: shade(table.cell(ri, ci), 'F1F5F9')
                doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if stripped.startswith('#'):
            level = min(len(stripped) - len(stripped.lstrip('#')), 3)
            heading = stripped[level:].strip()
            if level == 1 and heading == lines[0].lstrip('# ').strip():
                i += 1
                continue
            doc.add_heading(re.sub(r'[`*]', '', heading), level=level)
            i += 1
            continue
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s', '', stripped)
        else:
            p = doc.add_paragraph()
            text = stripped
        text = re.sub(r'\[([^\]]+)\]\([^)]*\)', r'\1', text)
        text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
        text = re.sub(r'[`*_]', '', text)
        p.add_run(text)
        i += 1

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run('VeraLane · 竞赛原型 · 仅供受控演示 · 2026-10-09')
    fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(110, 124, 140)
    out = ROOT / 'docs' / output_name
    doc.save(out)
    print(out)


build('安全自评与权限分级.md', 'VeraLane安全自评报告.docx', '安全自评报告')
build('VeraLane技术文档.md', 'VeraLane技术文档.docx', '技术文档')
build('部署说明.md', 'VeraLane部署说明.docx', '部署说明')
build('答辩讲稿.md', 'VeraLane答辩讲稿.docx', '答辩讲稿')
